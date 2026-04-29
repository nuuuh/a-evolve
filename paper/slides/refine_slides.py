#!/usr/bin/env python3
"""Refine the Self-Evolving Multi-Agent Orchestration presentation."""

import os
import copy
from pathlib import Path
from lxml import etree

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Circle
import matplotlib.patheffects as pe
import numpy as np
from PIL import Image

from pptx import Presentation
from pptx.util import Inches, Emu, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

# ── Constants ──────────────────────────────────────────────────────────

SLIDES_DIR = Path(__file__).parent
ASSETS_DIR = SLIDES_DIR / 'assets' / 'rendered'
PPTX_PATH = SLIDES_DIR / 'self_evolving_multi_agent_orchestration.pptx'

# Colors (matching existing palette)
ORANGE = '#C8571B'
PURPLE = '#6B5B95'
TEAL = '#2A9D8F'
NAVY = '#1B2A4A'
DARK_GREEN = '#2D6A4F'
LIGHT_FILL = '#EDF2F8'
WHITE = '#FFFFFF'
GRAY = '#888888'
RED = '#CC3333'
GREEN_CHECK = '#2D8F4E'

# Standard figure size matching existing custom diagrams
FIG_W, FIG_H = 14.56, 8.40
FIG_DPI = 100

# Standard slide shape positions (EMU)
TITLE_LEFT = 502920
TITLE_TOP = 219456
TITLE_W = 11201400
TITLE_H = 502920
SEP_LEFT = 502920
SEP_TOP = 786384
SEP_W = 11155680
SEP_H = 18288
IMG_LEFT = 1689100
IMG_TOP = 960120
IMG_W = 8810752
IMG_H = 4956048
FOOTER_LEFT = 685800
FOOTER_TOP = 6053328
FOOTER_W = 10789920
FOOTER_H = 301752
SOURCE_LEFT = 502920
SOURCE_TOP = 6556248
SOURCE_W = 11247120
SOURCE_H = 201168


def setup_style():
    plt.rcParams.update({
        'font.family': 'sans-serif',
        'font.sans-serif': ['DejaVu Sans', 'Arial', 'Helvetica'],
        'axes.facecolor': 'white',
        'figure.facecolor': 'white',
        'savefig.facecolor': 'white',
        'savefig.edgecolor': 'white',
        'savefig.bbox': 'tight',
        'savefig.pad_inches': 0.1,
    })


def save_fig(fig, name):
    path = ASSETS_DIR / name
    fig.savefig(str(path), dpi=FIG_DPI, bbox_inches='tight', pad_inches=0.3,
                facecolor='white', edgecolor='white')
    plt.close(fig)
    print(f'  Saved {path.name} ({os.path.getsize(path)} bytes)')
    return str(path)


def draw_rounded_box(ax, x, y, w, h, text, color=NAVY, fill=LIGHT_FILL,
                     fontsize=14, fontcolor=None, alpha=1.0, linewidth=2.5):
    box = FancyBboxPatch((x - w/2, y - h/2), w, h,
                         boxstyle="round,pad=0.02",
                         facecolor=fill, edgecolor=color,
                         linewidth=linewidth, alpha=alpha, zorder=2)
    ax.add_patch(box)
    fc = fontcolor or color
    ax.text(x, y, text, ha='center', va='center', fontsize=fontsize,
            color=fc, fontweight='bold', zorder=3)


def draw_arrow(ax, x1, y1, x2, y2, color=GRAY, linewidth=2, style='->', head_width=0.02):
    ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                arrowprops=dict(arrowstyle=style, color=color, lw=linewidth),
                zorder=1)


def draw_circle_node(ax, x, y, r, text, edgecolor=NAVY, fill=LIGHT_FILL,
                     fontsize=11, fontcolor=None, linewidth=2.5, lines=None):
    circle = Circle((x, y), r, facecolor=fill, edgecolor=edgecolor,
                    linewidth=linewidth, zorder=2)
    ax.add_patch(circle)
    fc = fontcolor or edgecolor
    if lines:
        for i, line in enumerate(lines):
            offset = (i - (len(lines)-1)/2) * fontsize * 0.015
            ax.text(x, y - offset, line, ha='center', va='center',
                    fontsize=fontsize, color=fc, fontweight='bold', zorder=3)
    else:
        ax.text(x, y, text, ha='center', va='center',
                fontsize=fontsize, color=fc, fontweight='bold', zorder=3)


# ═══════════════════════════════════════════════════════════════════════
# PHASE 1: Challenge diagram generation
# ═══════════════════════════════════════════════════════════════════════

def gen_c1_diagram():
    """Challenge C1: Static Topology -- two-panel comparison."""
    fig, (ax_l, ax_r) = plt.subplots(1, 2, figsize=(FIG_W, FIG_H))
    for ax in (ax_l, ax_r):
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.set_aspect('equal')
        ax.axis('off')

    # ── Left panel: The Problem ──
    ax_l.text(0.5, 0.95, 'The Problem', ha='center', va='top',
              fontsize=22, color=ORANGE, fontweight='bold')

    # Static chain
    chain_y = 0.72
    labels = ['Planner', 'Solver', 'Critic', 'Final']
    xs = [0.12, 0.35, 0.58, 0.81]
    bw, bh = 0.18, 0.10
    for i, (x, lbl) in enumerate(zip(xs, labels)):
        draw_rounded_box(ax_l, x, chain_y, bw, bh, lbl, color=NAVY,
                        fill=LIGHT_FILL, fontsize=11)
        if i < len(xs) - 1:
            draw_arrow(ax_l, x + bw/2 + 0.01, chain_y,
                      xs[i+1] - bw/2 - 0.01, chain_y, color=GRAY)

    ax_l.text(0.5, 0.60, 'One fixed pipeline for all tasks',
              ha='center', va='center', fontsize=13, color=GRAY, style='italic')

    # Task types that need different workflows
    task_y = 0.38
    tasks = [('Math\nproof', ORANGE), ('Web\nresearch', TEAL), ('Noisy\nevidence', DARK_GREEN)]
    task_xs = [0.2, 0.5, 0.8]
    for x, (label, color) in zip(task_xs, tasks):
        draw_circle_node(ax_l, x, task_y, 0.08, label, edgecolor=color,
                        fill=LIGHT_FILL, fontsize=10, fontcolor=color)

    # Big red X
    ax_l.text(0.5, 0.52, '✗', ha='center', va='center',
              fontsize=50, color=RED, fontweight='bold', zorder=5)

    ax_l.text(0.5, 0.18, 'Different tasks need\ndifferent structures',
              ha='center', va='center', fontsize=14, color=NAVY, fontweight='bold')

    # ── Right panel: What's Needed ──
    ax_r.text(0.5, 0.95, "What's Needed", ha='center', va='top',
              fontsize=22, color=TEAL, fontweight='bold')

    # Task-specific topologies
    # Math proof → chain
    ty = 0.72
    ax_r.text(0.15, ty, 'Math', ha='center', va='center',
              fontsize=13, color=ORANGE, fontweight='bold')
    for i, x in enumerate([0.35, 0.50, 0.65]):
        draw_circle_node(ax_r, x, ty, 0.035, '', edgecolor=ORANGE, fill='#FDE8D8')
        if i < 2:
            draw_arrow(ax_r, x + 0.04, ty, x + 0.11, ty, color=ORANGE, linewidth=1.5)
    ax_r.text(0.82, ty, 'chain', ha='center', va='center',
              fontsize=11, color=GRAY, style='italic')

    # Web research → tree
    ty = 0.50
    ax_r.text(0.15, ty, 'Search', ha='center', va='center',
              fontsize=13, color=TEAL, fontweight='bold')
    draw_circle_node(ax_r, 0.45, ty + 0.06, 0.035, '', edgecolor=TEAL, fill='#D4EDEA')
    for x in [0.35, 0.45, 0.55]:
        draw_circle_node(ax_r, x, ty - 0.06, 0.030, '', edgecolor=TEAL, fill='#D4EDEA')
        draw_arrow(ax_r, 0.45, ty + 0.02, x, ty - 0.03, color=TEAL, linewidth=1.5)
    ax_r.text(0.82, ty, 'tree', ha='center', va='center',
              fontsize=11, color=GRAY, style='italic')

    # Noisy evidence → convergent
    ty = 0.28
    ax_r.text(0.15, ty, 'Noisy', ha='center', va='center',
              fontsize=13, color=DARK_GREEN, fontweight='bold')
    for x in [0.35, 0.55]:
        draw_circle_node(ax_r, x, ty + 0.06, 0.030, '', edgecolor=DARK_GREEN, fill='#D4EDD8')
        draw_arrow(ax_r, x, ty + 0.03, 0.45, ty - 0.03, color=DARK_GREEN, linewidth=1.5)
    draw_circle_node(ax_r, 0.45, ty - 0.06, 0.035, '', edgecolor=DARK_GREEN, fill='#D4EDD8')
    ax_r.text(0.82, ty, 'graph', ha='center', va='center',
              fontsize=11, color=GRAY, style='italic')

    # Green checkmark
    ax_r.text(0.70, 0.50, '✓', ha='center', va='center',
              fontsize=50, color=GREEN_CHECK, fontweight='bold', zorder=5, alpha=0.3)

    ax_r.text(0.5, 0.10, 'Task-adaptive orchestration\nmatches structure to problem',
              ha='center', va='center', fontsize=14, color=NAVY, fontweight='bold')

    fig.tight_layout(pad=1.0)
    return save_fig(fig, 'custom_06_c1_v2.png')


def gen_c2_diagram():
    """Challenge C2: Supervision / Search Cost -- three-column comparison."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    ax.text(0.5, 0.96, 'How to learn good orchestration?',
            ha='center', va='top', fontsize=22, color=PURPLE, fontweight='bold')

    cols = [
        {'x': 0.18, 'label': 'RL Training', 'sublabel': '(Paper 1 approach)',
         'cost': 0.85, 'color': ORANGE, 'icon': '⟳', 'detail': 'Train policy on\ntask rewards'},
        {'x': 0.50, 'label': 'Validation Search', 'sublabel': '(ADAS, AFlow, MaAS)',
         'cost': 0.55, 'color': PURPLE, 'icon': '☐', 'detail': 'Tune MAS on\nlabeled val set'},
        {'x': 0.82, 'label': 'Zero Supervision', 'sublabel': '(the goal)',
         'cost': 0.15, 'color': TEAL, 'icon': '→', 'detail': 'Design MAS at\ninference time'},
    ]

    bar_bottom = 0.28
    bar_width = 0.12
    max_bar_h = 0.45

    for col in cols:
        x = col['x']
        c = col['color']
        h = col['cost'] * max_bar_h

        # Method label
        ax.text(x, 0.85, col['label'], ha='center', va='center',
                fontsize=16, color=c, fontweight='bold')
        ax.text(x, 0.79, col['sublabel'], ha='center', va='center',
                fontsize=11, color=GRAY)

        # Cost bar
        bar = FancyBboxPatch((x - bar_width/2, bar_bottom), bar_width, h,
                             boxstyle="round,pad=0.005",
                             facecolor=c, edgecolor=c, alpha=0.25, linewidth=0)
        ax.add_patch(bar)
        bar_border = FancyBboxPatch((x - bar_width/2, bar_bottom), bar_width, h,
                                    boxstyle="round,pad=0.005",
                                    facecolor='none', edgecolor=c, linewidth=2.5)
        ax.add_patch(bar_border)

        # Cost label on bar
        cost_labels = {0.85: 'HIGH', 0.55: 'MEDIUM', 0.15: 'LOW'}
        ax.text(x, bar_bottom + h + 0.03, cost_labels.get(col['cost'], ''),
                ha='center', va='bottom', fontsize=14, color=c, fontweight='bold')

        # Detail below bar
        ax.text(x, bar_bottom - 0.04, col['detail'], ha='center', va='top',
                fontsize=11, color=NAVY)

    # Y-axis label
    ax.text(0.04, bar_bottom + max_bar_h / 2, 'Supervision\ncost',
            ha='center', va='center', fontsize=13, color=GRAY, rotation=90)

    # Bottom arrow
    ax.annotate('', xy=(0.88, 0.10), xytext=(0.12, 0.10),
                arrowprops=dict(arrowstyle='->', color=TEAL, lw=3))
    ax.text(0.50, 0.06, 'Decreasing supervision requirement',
            ha='center', va='center', fontsize=14, color=TEAL, fontweight='bold')

    return save_fig(fig, 'custom_07_c2_v2.png')


def gen_c3_diagram():
    """Challenge C3: When Does MAS Help? -- performance-cost scatter."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Main plot area
    px, py, pw, ph = 0.08, 0.12, 0.55, 0.75
    plot_rect = FancyBboxPatch((px, py), pw, ph,
                               boxstyle="round,pad=0.01",
                               facecolor=WHITE, edgecolor=NAVY,
                               linewidth=1.5, zorder=0)
    ax.add_patch(plot_rect)

    # Axes
    ax.annotate('', xy=(px + pw - 0.02, py + 0.01),
                xytext=(px + 0.02, py + 0.01),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=1.5))
    ax.annotate('', xy=(px + 0.02, py + ph - 0.02),
                xytext=(px + 0.02, py + 0.02),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=1.5))
    ax.text(px + pw / 2, py - 0.03, 'Cost (tokens)', ha='center', fontsize=13, color=NAVY)
    ax.text(px - 0.02, py + ph / 2, 'Accuracy', ha='center', va='center',
            fontsize=13, color=NAVY, rotation=90)

    # Points
    sa_x, sa_y = px + 0.12, py + 0.30
    naive_x, naive_y = px + 0.42, py + 0.35
    smart_x, smart_y = px + 0.25, py + 0.58

    # Single Agent
    ax.plot(sa_x, sa_y, 'o', markersize=18, color=NAVY, zorder=4)
    ax.text(sa_x + 0.02, sa_y - 0.06, 'Single\nAgent', ha='center',
            fontsize=11, color=NAVY, fontweight='bold')

    # Naive MAS
    ax.plot(naive_x, naive_y, 'o', markersize=18, color=ORANGE, zorder=4)
    ax.text(naive_x + 0.02, naive_y - 0.06, 'Naive\nMAS', ha='center',
            fontsize=11, color=ORANGE, fontweight='bold')

    # Smart MAS
    ax.plot(smart_x, smart_y, '*', markersize=28, color=TEAL, zorder=4)
    ax.text(smart_x + 0.06, smart_y + 0.02, 'Smart\nMAS', ha='center',
            fontsize=11, color=TEAL, fontweight='bold')

    # Annotations
    ax.annotate('more agents,\nsame accuracy,\nhigher cost',
                xy=(naive_x - 0.02, naive_y + 0.02),
                xytext=(sa_x + 0.10, sa_y + 0.15),
                fontsize=10, color=RED,
                arrowprops=dict(arrowstyle='->', color=RED, lw=1.5, connectionstyle='arc3,rad=0.2'))

    ax.annotate('smarter\norchestration',
                xy=(smart_x + 0.02, smart_y - 0.02),
                xytext=(naive_x - 0.05, naive_y + 0.18),
                fontsize=10, color=GREEN_CHECK,
                arrowprops=dict(arrowstyle='->', color=GREEN_CHECK, lw=1.5, connectionstyle='arc3,rad=-0.2'))

    # Pareto frontier hint
    pf_x = np.array([px + 0.08, smart_x, px + pw - 0.08])
    pf_y = np.array([py + 0.15, smart_y + 0.05, py + ph - 0.10])
    ax.plot(pf_x, pf_y, '--', color=TEAL, alpha=0.4, linewidth=2, zorder=1)
    ax.text(px + pw - 0.10, py + ph - 0.06, 'Pareto\nfrontier', ha='center',
            fontsize=10, color=TEAL, alpha=0.6, style='italic')

    # Right callout box
    bx, by, bw_r, bh_r = 0.68, 0.20, 0.28, 0.55
    callout = FancyBboxPatch((bx, by), bw_r, bh_r,
                             boxstyle="round,pad=0.02",
                             facecolor='#E8F5F3', edgecolor=TEAL,
                             linewidth=2.5, zorder=2)
    ax.add_patch(callout)

    ax.text(bx + bw_r/2, by + bh_r - 0.04, 'When does MAS help?',
            ha='center', va='top', fontsize=14, color=TEAL, fontweight='bold')

    factors = [
        'Task is decomposable',
        'Sub-agents at edge of\ncompetence',
        'Adversarial or noisy\ninputs present',
        'Parallel sub-tasks exist',
    ]
    fy = by + bh_r - 0.13
    for f in factors:
        ax.text(bx + 0.04, fy, '•  ' + f, ha='left', va='top',
                fontsize=11, color=NAVY)
        fy -= 0.11

    ax.text(bx + bw_r/2, by + 0.03, '→ Paper 3 answers this\n    systematically',
            ha='center', va='bottom', fontsize=11, color=TEAL, fontweight='bold')

    return save_fig(fig, 'custom_08_c3_v2.png')


# ═══════════════════════════════════════════════════════════════════════
# PHASE 2: Fix duplicate content diagrams
# ═══════════════════════════════════════════════════════════════════════

def gen_p1_motivation():
    """Paper 1 Motivation: static vs dynamic comparison."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Left: faded static chain
    ax.text(0.22, 0.90, 'Current: Static MAS', ha='center', va='top',
            fontsize=18, color=GRAY, fontweight='bold')

    chain_y = 0.65
    labels = ['Planner', 'Solver', 'Critic', 'Final']
    xs = [0.06, 0.17, 0.28, 0.39]
    for i, (x, lbl) in enumerate(zip(xs, labels)):
        draw_rounded_box(ax, x, chain_y, 0.10, 0.08, lbl, color=GRAY,
                        fill='#F0F0F0', fontsize=9, alpha=0.6)
        if i < len(xs) - 1:
            draw_arrow(ax, x + 0.06, chain_y, xs[i+1] - 0.06, chain_y,
                      color=GRAY, linewidth=1.5)

    # Red strikethrough
    ax.plot([0.02, 0.42], [chain_y + 0.06, chain_y - 0.06],
            color=RED, linewidth=3, alpha=0.7, zorder=5)
    ax.plot([0.02, 0.42], [chain_y - 0.06, chain_y + 0.06],
            color=RED, linewidth=3, alpha=0.7, zorder=5)

    ax.text(0.22, 0.42, 'Fixed roles, fixed path,\nfixed execution order',
            ha='center', va='center', fontsize=13, color=RED, style='italic')

    # Big arrow in the middle
    ax.annotate('', xy=(0.62, 0.60), xytext=(0.48, 0.60),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=4))
    ax.text(0.55, 0.52, "Paper 1's\ncontribution", ha='center', va='center',
            fontsize=12, color=NAVY, fontweight='bold')

    # Right: dynamic orchestration
    ax.text(0.78, 0.90, 'Proposed: Dynamic Orchestration', ha='center', va='top',
            fontsize=18, color=ORANGE, fontweight='bold')

    # Central orchestrator
    draw_rounded_box(ax, 0.78, 0.62, 0.16, 0.10, 'Puppeteer\n(learned policy)',
                    color=NAVY, fill=LIGHT_FILL, fontsize=11)

    # Surrounding agents
    agents = [
        (0.62, 0.80, 'Planner', ORANGE),
        (0.94, 0.80, 'Searcher', TEAL),
        (0.62, 0.44, 'Solver', DARK_GREEN),
        (0.94, 0.44, 'Critic', PURPLE),
    ]
    for ax_x, ay, label, color in agents:
        draw_rounded_box(ax, ax_x, ay, 0.12, 0.08, label, color=color,
                        fill=LIGHT_FILL, fontsize=10)

    # Dynamic arrows (dashed to show adaptive routing)
    for ax_x, ay, _, color in agents:
        dx = ax_x - 0.78
        dy = ay - 0.62
        norm = (dx**2 + dy**2)**0.5
        start_x = 0.78 + dx/norm * 0.09
        start_y = 0.62 + dy/norm * 0.06
        end_x = ax_x - dx/norm * 0.07
        end_y = ay - dy/norm * 0.05
        ax.annotate('', xy=(end_x, end_y), xytext=(start_x, start_y),
                    arrowprops=dict(arrowstyle='->', color=color, lw=2,
                                   linestyle='dashed'))

    ax.text(0.78, 0.30, 'Agents activated dynamically\nbased on task state',
            ha='center', va='center', fontsize=13, color=NAVY, fontweight='bold')

    # Bottom
    ax.text(0.50, 0.12, 'Key insight: the orchestration policy — not the agents — should evolve',
            ha='center', va='center', fontsize=15, color=NAVY,
            fontweight='bold', style='italic',
            bbox=dict(boxstyle='round,pad=0.4', facecolor=LIGHT_FILL, edgecolor=NAVY, linewidth=1.5))

    return save_fig(fig, 'custom_10_p1_motivation_v2.png')


def gen_meta_loop_unrolled():
    """MAS-ZERO meta-agent loop unrolled over iterations (slide 23)."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Three rows: MAS-Init, MAS-Evolve iterations, MAS-Verify
    rows = [
        {'y': 0.82, 'label': 'Step 1: MAS-Init', 'color': NAVY,
         'items': ['CoT', 'CoT-SC', 'Debate', 'Self-Refine'],
         'desc': 'Run building blocks → initial candidates'},
        {'y': 0.52, 'label': 'Step 2: MAS-Evolve (×T iterations)', 'color': ORANGE,
         'items': ['Meta-Design', 'Execute\nMAS', 'Meta-\nFeedback', 'Experience\nLibrary'],
         'desc': 'Iteratively design → run → critique → store experience'},
        {'y': 0.22, 'label': 'Step 3: MAS-Verify', 'color': TEAL,
         'items': ['Rank by\nfrequency', 'Filter\ninvalid', 'Select\nbest'],
         'desc': 'Choose most reliable candidate answer'},
    ]

    for row in rows:
        y = row['y']
        c = row['color']

        # Row label
        ax.text(0.02, y + 0.08, row['label'], ha='left', va='center',
                fontsize=15, color=c, fontweight='bold')

        # Items as boxes
        items = row['items']
        n = len(items)
        x_start = 0.15
        x_end = 0.88
        spacing = (x_end - x_start) / n
        for i, item in enumerate(items):
            x = x_start + spacing * (i + 0.5)
            draw_rounded_box(ax, x, y, spacing * 0.75, 0.10, item,
                           color=c, fill=LIGHT_FILL, fontsize=11)
            if i < n - 1:
                draw_arrow(ax, x + spacing * 0.40, y,
                          x + spacing * 0.60, y, color=c, linewidth=2)

        # Description
        ax.text(0.52, y - 0.08, row['desc'], ha='center', va='center',
                fontsize=12, color=GRAY, style='italic')

    # Vertical arrows between rows
    for y1, y2 in [(0.72, 0.64), (0.42, 0.34)]:
        ax.annotate('', xy=(0.52, y2), xytext=(0.52, y1),
                    arrowprops=dict(arrowstyle='->', color=NAVY, lw=2.5))

    # Iteration loop arrow on the right
    ax.annotate('', xy=(0.95, 0.60), xytext=(0.95, 0.44),
                arrowprops=dict(arrowstyle='->', color=ORANGE, lw=2.5,
                               connectionstyle='arc3,rad=-0.5'))
    ax.text(0.97, 0.52, 'refine', ha='left', va='center',
            fontsize=11, color=ORANGE, style='italic', rotation=90)

    return save_fig(fig, 'custom_23_meta_loop_v2.png')


# ═══════════════════════════════════════════════════════════════════════
# PHASE 3: Paper 3 expansion diagrams
# ═══════════════════════════════════════════════════════════════════════

def gen_what_evolves_p3():
    """What Evolves in MAS-Orchestra: holistic single-step generation."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Left: Sequential decisions (Papers 1 & 2)
    ax.text(0.22, 0.93, 'Sequential Decisions', ha='center', va='top',
            fontsize=18, color=GRAY, fontweight='bold')
    ax.text(0.22, 0.86, '(Papers 1 & 2)', ha='center', va='top',
            fontsize=13, color=GRAY)

    steps = ['Pick\nagent 1', 'Pick\nagent 2', 'Pick\nagent 3', '...']
    for i, (label) in enumerate(steps):
        y = 0.72 - i * 0.14
        draw_rounded_box(ax, 0.22, y, 0.18, 0.09, label, color=GRAY,
                        fill='#F0F0F0', fontsize=11, fontcolor=GRAY)
        if i < len(steps) - 1:
            draw_arrow(ax, 0.22, y - 0.05, 0.22, y - 0.09, color=GRAY)

    ax.text(0.22, 0.22, 'Step-by-step,\nlocal decisions', ha='center', va='center',
            fontsize=13, color=GRAY, style='italic')

    # Arrow
    ax.annotate('', xy=(0.54, 0.55), xytext=(0.40, 0.55),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=4))

    # Right: Holistic orchestration (Paper 3)
    ax.text(0.75, 0.93, 'Holistic Orchestration', ha='center', va='top',
            fontsize=18, color=TEAL, fontweight='bold')
    ax.text(0.75, 0.86, '(Paper 3: MAS-Orchestra)', ha='center', va='top',
            fontsize=13, color=TEAL)

    # Orchestrator box
    draw_rounded_box(ax, 0.75, 0.72, 0.22, 0.10, 'Orchestrator', color=NAVY,
                    fill=LIGHT_FILL, fontsize=14)

    # Function calls bracket
    funcs = [
        ('create_agent("Solver")', ORANGE),
        ('create_agent("Searcher")', TEAL),
        ('create_flow(Solver→Searcher)', DARK_GREEN),
    ]
    for i, (func, color) in enumerate(funcs):
        y = 0.55 - i * 0.10
        draw_rounded_box(ax, 0.75, y, 0.30, 0.07, func, color=color,
                        fill=LIGHT_FILL, fontsize=10)

    ax.annotate('', xy=(0.75, 0.58), xytext=(0.75, 0.65),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=2))
    ax.text(0.58, 0.45, '} single\n  step', ha='center', va='center',
            fontsize=16, color=NAVY, fontweight='bold')

    # DoM box at bottom
    dom_y = 0.18
    ax.text(0.75, dom_y + 0.08, 'Degree of MAS (DoM)', ha='center', va='center',
            fontsize=14, color=NAVY, fontweight='bold')
    # Low DoM
    draw_rounded_box(ax, 0.60, dom_y - 0.04, 0.16, 0.07, 'Low: ≤1 agent',
                    color=PURPLE, fill=LIGHT_FILL, fontsize=10)
    # High DoM
    draw_rounded_box(ax, 0.90, dom_y - 0.04, 0.16, 0.07, 'High: full MAS',
                    color=TEAL, fill=LIGHT_FILL, fontsize=10)
    ax.annotate('', xy=(0.82, dom_y - 0.04), xytext=(0.68, dom_y - 0.04),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=2))

    return save_fig(fig, 'custom_31b_what_evolves_p3.png')


def gen_masbench_axes():
    """MASBench: Five Axes of Complexity."""
    fig, axes = plt.subplots(1, 5, figsize=(FIG_W, FIG_H))

    axis_info = [
        {'name': 'Depth', 'color': '#E07B39', 'icon': 'chain',
         'defn': 'Length of longest\ndependency chain',
         'verify': 'Final answer\nonly'},
        {'name': 'Horizon', 'color': '#C8571B', 'icon': 'horizon',
         'defn': 'Intermediate answers\ncarried forward',
         'verify': 'Intermediate\n+ final'},
        {'name': 'Breadth', 'color': '#6B5B95', 'icon': 'breadth',
         'defn': 'Max dependencies\nper sub-task',
         'verify': 'Final answer\nonly'},
        {'name': 'Parallel', 'color': '#2A9D8F', 'icon': 'parallel',
         'defn': 'Independent\nsub-task components',
         'verify': 'Final (multiple\nanswers)'},
        {'name': 'Robustness', 'color': '#2D6A4F', 'icon': 'robust',
         'defn': 'Sub-tasks with\nadversarial attacks',
         'verify': 'Intermediate\n+ adversarial'},
    ]

    for i, (ax, info) in enumerate(zip(axes, axis_info)):
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.set_aspect('auto')
        ax.axis('off')

        c = info['color']

        # Background card
        card = FancyBboxPatch((0.05, 0.05), 0.90, 0.90,
                              boxstyle="round,pad=0.03",
                              facecolor=WHITE, edgecolor=c,
                              linewidth=3, zorder=0)
        ax.add_patch(card)

        # Axis name
        ax.text(0.5, 0.92, info['name'], ha='center', va='top',
                fontsize=18, color=c, fontweight='bold')

        # Task structure icon
        icon_y = 0.68
        if info['icon'] == 'chain':
            for j in range(3):
                cx = 0.25 + j * 0.25
                draw_circle_node(ax, cx, icon_y, 0.06, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
                if j < 2:
                    draw_arrow(ax, cx + 0.07, icon_y, cx + 0.18, icon_y, color=c, linewidth=1.5)
        elif info['icon'] == 'horizon':
            positions = [(0.2, icon_y+0.06), (0.5, icon_y+0.06), (0.5, icon_y-0.06), (0.8, icon_y-0.06)]
            for j, (cx, cy) in enumerate(positions):
                draw_circle_node(ax, cx, cy, 0.05, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
                if j < len(positions) - 1:
                    nx, ny = positions[j+1]
                    draw_arrow(ax, cx + 0.06, cy, nx - 0.06, ny, color=c, linewidth=1.5)
        elif info['icon'] == 'breadth':
            draw_circle_node(ax, 0.5, icon_y + 0.06, 0.05, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
            for cx in [0.25, 0.5, 0.75]:
                draw_circle_node(ax, cx, icon_y - 0.06, 0.04, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
                draw_arrow(ax, 0.5, icon_y, cx, icon_y - 0.02, color=c, linewidth=1.5)
        elif info['icon'] == 'parallel':
            for cx in [0.25, 0.5, 0.75]:
                draw_circle_node(ax, cx, icon_y + 0.04, 0.04, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
                draw_circle_node(ax, cx, icon_y - 0.08, 0.04, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
                draw_arrow(ax, cx, icon_y - 0.01, cx, icon_y - 0.04, color=c, linewidth=1.5)
        elif info['icon'] == 'robust':
            for j, cx in enumerate([0.3, 0.7]):
                draw_circle_node(ax, cx, icon_y + 0.04, 0.05, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
                draw_circle_node(ax, cx, icon_y - 0.08, 0.04, '', edgecolor=c, fill=LIGHT_FILL, linewidth=2)
                draw_arrow(ax, cx, icon_y - 0.02, cx, icon_y - 0.04, color=c, linewidth=1.5)
            ax.text(0.7, icon_y + 0.04, '!', ha='center', va='center',
                    fontsize=14, color=RED, fontweight='bold', zorder=5)

        # Definition
        ax.text(0.5, 0.42, info['defn'], ha='center', va='center',
                fontsize=12, color=NAVY, fontweight='bold')

        # Verification type
        ax.text(0.5, 0.20, 'Verification:', ha='center', va='center',
                fontsize=10, color=GRAY)
        ax.text(0.5, 0.12, info['verify'], ha='center', va='center',
                fontsize=10, color=c, fontweight='bold')

    fig.tight_layout(pad=0.5)
    return save_fig(fig, 'custom_33b_masbench_axes.png')


def gen_edge_competence():
    """Key Insight: MAS is most effective at the edge of sub-agent competence."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Plot area
    px, py, pw, ph = 0.10, 0.15, 0.52, 0.70

    # X axis
    ax.annotate('', xy=(px + pw, py), xytext=(px, py),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=2))
    ax.text(px + pw / 2, py - 0.06, 'Sub-agent Competence', ha='center',
            fontsize=14, color=NAVY, fontweight='bold')

    # Y axis
    ax.annotate('', xy=(px, py + ph), xytext=(px, py),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=2))
    ax.text(px - 0.04, py + ph / 2, 'MAS Benefit\nover SAS', ha='center', va='center',
            fontsize=13, color=NAVY, fontweight='bold', rotation=90)

    # Inverted-U curve
    x_vals = np.linspace(0, 1, 100)
    y_vals = 4 * x_vals * (1 - x_vals)  # parabola peaking at x=0.5
    curve_x = px + 0.03 + x_vals * (pw - 0.06)
    curve_y = py + 0.03 + y_vals * (ph - 0.06)
    ax.plot(curve_x, curve_y, color=TEAL, linewidth=4, zorder=3)
    ax.fill_between(curve_x, py, curve_y, alpha=0.08, color=TEAL, zorder=1)

    # Three zones
    zone_w = (pw - 0.06) / 3
    zones = [
        {'x': px + 0.03 + zone_w * 0.5, 'label': 'Too Easy', 'sublabel': 'SAS suffices',
         'color': GRAY},
        {'x': px + 0.03 + zone_w * 1.5, 'label': 'Sweet Spot', 'sublabel': 'MAS shines',
         'color': TEAL},
        {'x': px + 0.03 + zone_w * 2.5, 'label': 'Too Hard', 'sublabel': 'Nothing helps',
         'color': GRAY},
    ]
    for z in zones:
        ax.text(z['x'], py + 0.03, z['label'], ha='center', va='bottom',
                fontsize=12, color=z['color'], fontweight='bold')
        ax.text(z['x'], py - 0.01, z['sublabel'], ha='center', va='top',
                fontsize=10, color=z['color'], style='italic')

    # Peak annotation
    peak_x = px + 0.03 + 0.5 * (pw - 0.06)
    peak_y = py + 0.03 + 1.0 * (ph - 0.06)
    ax.plot(peak_x, peak_y, '*', markersize=20, color=ORANGE, zorder=5)
    ax.text(peak_x + 0.04, peak_y + 0.01, 'Edge of\ncompetence',
            fontsize=12, color=ORANGE, fontweight='bold')

    # Right callout boxes
    findings = [
        ('Instruction-tuned LLMs\n> RLMs as orchestrators', PURPLE, 0.72),
        ('MAS uniquely robust\nto adversarial noise', TEAL, 0.52),
        ('MAS benefit depends on\ntask structure + DoM', ORANGE, 0.32),
    ]
    for text, color, fy in findings:
        box = FancyBboxPatch((0.67, fy - 0.06), 0.30, 0.13,
                             boxstyle="round,pad=0.015",
                             facecolor=WHITE, edgecolor=color,
                             linewidth=2.5, zorder=2)
        ax.add_patch(box)
        ax.text(0.82, fy + 0.005, text, ha='center', va='center',
                fontsize=11, color=NAVY)

    ax.text(0.82, 0.90, 'Key Findings', ha='center', va='center',
            fontsize=16, color=NAVY, fontweight='bold')

    return save_fig(fig, 'custom_34b_edge_competence.png')


# ═══════════════════════════════════════════════════════════════════════
# PHASE 4: Transition diagrams
# ═══════════════════════════════════════════════════════════════════════

def gen_transition_p1_p2():
    """Transition slide: Paper 1 → Paper 2."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Paper 1 box
    draw_rounded_box(ax, 0.22, 0.55, 0.30, 0.22, '', color=ORANGE, fill='#FDF0E8',
                    fontsize=1, linewidth=3)
    ax.text(0.22, 0.62, 'Paper 1: Puppeteer', ha='center', va='center',
            fontsize=16, color=ORANGE, fontweight='bold')
    ax.text(0.22, 0.52, 'Learned dynamic\norchestration via RL', ha='center', va='center',
            fontsize=13, color=NAVY)
    ax.text(0.22, 0.42, '✓ Solves C1 (static topology)', ha='center', va='center',
            fontsize=12, color=GREEN_CHECK, fontweight='bold')

    # Arrow
    ax.annotate('', xy=(0.60, 0.55), xytext=(0.40, 0.55),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=4))

    # Question
    ax.text(0.50, 0.75, 'But RL training requires reward signals,\ncompute, and labeled data (C2)...',
            ha='center', va='center', fontsize=14, color=PURPLE,
            fontweight='bold', style='italic',
            bbox=dict(boxstyle='round,pad=0.4', facecolor='#F0EBF5',
                     edgecolor=PURPLE, linewidth=2))

    # Paper 2 box
    draw_rounded_box(ax, 0.78, 0.55, 0.30, 0.22, '', color=PURPLE, fill='#F0EBF5',
                    fontsize=1, linewidth=3)
    ax.text(0.78, 0.62, 'Paper 2: MAS-ZERO', ha='center', va='center',
            fontsize=16, color=PURPLE, fontweight='bold')
    ax.text(0.78, 0.52, 'Self-evolved MAS design\nat inference time', ha='center', va='center',
            fontsize=13, color=NAVY)
    ax.text(0.78, 0.42, '→ Zero supervision needed', ha='center', va='center',
            fontsize=12, color=TEAL, fontweight='bold')

    # Bottom: narrative thread
    ax.text(0.50, 0.20, 'Can we achieve adaptive orchestration without any training?',
            ha='center', va='center', fontsize=18, color=NAVY, fontweight='bold')

    return save_fig(fig, 'custom_19b_transition.png')


def gen_transition_p2_p3():
    """Transition slide: Paper 2 → Paper 3."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Paper 2 box
    draw_rounded_box(ax, 0.22, 0.55, 0.30, 0.22, '', color=PURPLE, fill='#F0EBF5',
                    fontsize=1, linewidth=3)
    ax.text(0.22, 0.62, 'Paper 2: MAS-ZERO', ha='center', va='center',
            fontsize=16, color=PURPLE, fontweight='bold')
    ax.text(0.22, 0.52, 'Inference-time\nMAS self-design', ha='center', va='center',
            fontsize=13, color=NAVY)
    ax.text(0.22, 0.42, '✓ Solves C2 (no supervision)', ha='center', va='center',
            fontsize=12, color=GREEN_CHECK, fontweight='bold')

    # Arrow
    ax.annotate('', xy=(0.60, 0.55), xytext=(0.40, 0.55),
                arrowprops=dict(arrowstyle='->', color=NAVY, lw=4))

    # Question
    ax.text(0.50, 0.75, 'But when should we even use MAS?\nWhen does it help vs. add overhead? (C3)',
            ha='center', va='center', fontsize=14, color=TEAL,
            fontweight='bold', style='italic',
            bbox=dict(boxstyle='round,pad=0.4', facecolor='#E8F5F3',
                     edgecolor=TEAL, linewidth=2))

    # Paper 3 box
    draw_rounded_box(ax, 0.78, 0.55, 0.30, 0.22, '', color=TEAL, fill='#E8F5F3',
                    fontsize=1, linewidth=3)
    ax.text(0.78, 0.62, 'Paper 3: MAS-Orchestra', ha='center', va='center',
            fontsize=16, color=TEAL, fontweight='bold')
    ax.text(0.78, 0.52, 'Holistic orchestration +\ncontrolled benchmarks', ha='center', va='center',
            fontsize=13, color=NAVY)
    ax.text(0.78, 0.42, '→ Systematic MAS analysis', ha='center', va='center',
            fontsize=12, color=ORANGE, fontweight='bold')

    # Bottom
    ax.text(0.50, 0.20, 'When and why does multi-agent coordination actually help?',
            ha='center', va='center', fontsize=18, color=NAVY, fontweight='bold')

    return save_fig(fig, 'custom_29b_transition.png')


# ═══════════════════════════════════════════════════════════════════════
# PHASE 5: Fix slide 3 text overflow
# ═══════════════════════════════════════════════════════════════════════

def gen_five_decisions_fixed():
    """Regenerate the 5 orchestration decisions with fixed text wrapping."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.set_aspect('equal')
    ax.axis('off')

    # Pentagon layout
    cx, cy = 0.5, 0.48
    r = 0.30
    decisions = [
        ('Role\nassignment', PURPLE, 90),
        ('Communication\ntopology', TEAL, 162),
        ('Execution\norder', DARK_GREEN, 234),
        ('Verification\naggregation', NAVY, 306),
        ('Agent\nselection', ORANGE, 18),
    ]

    positions = []
    for label, color, angle_deg in decisions:
        angle = np.radians(angle_deg)
        x = cx + r * np.cos(angle)
        y = cy + r * np.sin(angle)
        positions.append((x, y))

    # Draw connecting lines (pentagon edges + diagonals)
    for i in range(len(positions)):
        for j in range(i + 1, len(positions)):
            x1, y1 = positions[i]
            x2, y2 = positions[j]
            ax.plot([x1, x2], [y1, y2], color=LIGHT_FILL, linewidth=2.5, zorder=0)

    # Draw circles with labels
    for (x, y), (label, color, _) in zip(positions, decisions):
        circle_r = 0.11
        circle = Circle((x, y), circle_r, facecolor=LIGHT_FILL, edgecolor=color,
                        linewidth=3, zorder=2)
        ax.add_patch(circle)
        ax.text(x, y, label, ha='center', va='center',
                fontsize=13, color=color, fontweight='bold', zorder=3)

    return save_fig(fig, 'custom_03_five_decisions_v2.png')


def gen_final_synthesis():
    """Updated final synthesis table with Key Mechanism column."""
    fig, ax = plt.subplots(1, 1, figsize=(FIG_W, FIG_H))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    # Table
    headers = ['Paper', 'Key Mechanism', 'Orchestration\nContribution', 'Challenge']
    colors_row = [ORANGE, PURPLE, TEAL]
    rows = [
        ['Evolving\nOrchestration', 'RL policy\n(REINFORCE)', 'dynamic\ncontroller', 'C1'],
        ['MAS-ZERO', 'meta-agent\nloop', 'test-time\nself-design', 'C2'],
        ['MAS-Orchestra', 'function-calling\nRL + DoM', 'when MAS\nhelps', 'C3'],
    ]

    n_cols = len(headers)
    n_rows = len(rows) + 1  # +1 for header
    col_w = 0.22
    row_h = 0.14
    table_w = col_w * n_cols
    table_x = (1 - table_w) / 2
    table_y = 0.82

    # Header row
    for j, header in enumerate(headers):
        x = table_x + j * col_w
        box = FancyBboxPatch((x + 0.005, table_y - row_h + 0.005), col_w - 0.01, row_h - 0.01,
                             boxstyle="round,pad=0.005",
                             facecolor=LIGHT_FILL, edgecolor=NAVY, linewidth=2)
        ax.add_patch(box)
        ax.text(x + col_w/2, table_y - row_h/2, header, ha='center', va='center',
                fontsize=13, color=NAVY, fontweight='bold')

    # Data rows
    for i, (row, c) in enumerate(zip(rows, colors_row)):
        y = table_y - (i + 1) * row_h
        for j, cell in enumerate(row):
            x = table_x + j * col_w
            box = FancyBboxPatch((x + 0.005, y - row_h + 0.005), col_w - 0.01, row_h - 0.01,
                                 boxstyle="round,pad=0.005",
                                 facecolor=WHITE, edgecolor='#DDDDDD', linewidth=1)
            ax.add_patch(box)
            fc = c if j == 3 else NAVY
            ax.text(x + col_w/2, y - row_h/2, cell, ha='center', va='center',
                    fontsize=12, color=fc, fontweight='bold' if j == 3 else 'normal')

    # Footer text
    ax.text(0.5, 0.18, 'Next step: navigate and preserve evolved strategies over time.',
            ha='center', va='center', fontsize=16, color=NAVY, fontweight='bold')

    return save_fig(fig, 'custom_35_final_synthesis_v2.png')


# ═══════════════════════════════════════════════════════════════════════
# PPTX MANIPULATION
# ═══════════════════════════════════════════════════════════════════════

def replace_slide_image(slide, new_image_path):
    """Replace the first PICTURE shape on a slide with a new image."""
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            sp = shape._element
            sp.getparent().remove(sp)
            slide.shapes.add_picture(new_image_path, left, top, width, height)
            return True
    return False


def update_title_text(slide, new_text):
    """Update the first TextBox (title) text, preserving formatting."""
    shape = slide.shapes[0]
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            if para.runs:
                full_text = ''.join(r.text for r in para.runs)
                para.runs[0].text = new_text
                for r in para.runs[1:]:
                    r.text = ''
                return True
    return False


def update_footer_text(slide, new_text, shape_idx=3):
    """Update footer TextBox text."""
    if shape_idx < len(slide.shapes):
        shape = slide.shapes[shape_idx]
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                if para.runs:
                    para.runs[0].text = new_text
                    for r in para.runs[1:]:
                        r.text = ''
                    return True
    return False


def insert_slide_after(prs, after_idx, title_text, image_path, footer_text, source_text=None):
    """Insert a new content slide after the given index."""
    # Use Blank layout (index 6 in most templates)
    blank_layout = None
    for layout in prs.slide_layouts:
        if layout.name == 'Blank':
            blank_layout = layout
            break
    if blank_layout is None:
        blank_layout = prs.slide_layouts[6]

    new_slide = prs.slides.add_slide(blank_layout)

    # Move to correct position
    slide_list = prs.slides._sldIdLst
    sld_ids = list(slide_list)
    new_elem = sld_ids[-1]
    slide_list.remove(new_elem)
    target = sld_ids[after_idx]
    target.addnext(new_elem)

    # Add title TextBox
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor

    txBox = new_slide.shapes.add_textbox(Emu(TITLE_LEFT), Emu(TITLE_TOP),
                                         Emu(TITLE_W), Emu(TITLE_H))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title_text
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1B, 0x2A, 0x4A)
    run.font.bold = True

    # Add separator line
    from pptx.enum.shapes import MSO_SHAPE
    sep = new_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Emu(SEP_LEFT), Emu(SEP_TOP),
                                     Emu(SEP_W), Emu(SEP_H))
    sep.fill.solid()
    sep.fill.fore_color.rgb = RGBColor(0x1B, 0x2A, 0x4A)
    sep.line.fill.background()

    # Add image
    new_slide.shapes.add_picture(image_path, Emu(IMG_LEFT), Emu(IMG_TOP),
                                 Emu(IMG_W), Emu(IMG_H))

    # Add footer TextBox
    ftBox = new_slide.shapes.add_textbox(Emu(FOOTER_LEFT), Emu(FOOTER_TOP),
                                         Emu(FOOTER_W), Emu(FOOTER_H))
    ft_tf = ftBox.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.alignment = PP_ALIGN.CENTER
    ft_run = ft_p.add_run()
    ft_run.text = footer_text
    ft_run.font.size = Pt(14)
    ft_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # Add source citation if provided
    if source_text:
        srcBox = new_slide.shapes.add_textbox(Emu(SOURCE_LEFT), Emu(SOURCE_TOP),
                                              Emu(SOURCE_W), Emu(SOURCE_H))
        src_tf = srcBox.text_frame
        src_p = src_tf.paragraphs[0]
        src_run = src_p.add_run()
        src_run.text = source_text
        src_run.font.size = Pt(10)
        src_run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

    return new_slide


def crop_paper_image(src_name, crop_bottom_pct=0.18):
    """Crop the bottom caption from a paper screenshot."""
    src_path = ASSETS_DIR / src_name
    if not src_path.exists():
        print(f'  WARNING: {src_name} not found, skipping crop')
        return None

    img = Image.open(src_path)
    w, h = img.size
    crop_h = int(h * (1 - crop_bottom_pct))
    cropped = img.crop((0, 0, w, crop_h))

    out_name = src_name.replace('.png', '_cropped.png')
    out_path = ASSETS_DIR / out_name
    cropped.save(str(out_path))
    print(f'  Cropped {src_name} → {out_name} (removed bottom {crop_bottom_pct*100:.0f}%)')
    return str(out_path)


# ═══════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════

def main():
    setup_style()

    print('=== Phase 1: Generating challenge diagrams ===')
    c1_path = gen_c1_diagram()
    c2_path = gen_c2_diagram()
    c3_path = gen_c3_diagram()

    print('\n=== Phase 2: Generating duplicate-fix diagrams ===')
    p1_motiv_path = gen_p1_motivation()
    meta_loop_path = gen_meta_loop_unrolled()

    print('\n=== Phase 3: Generating Paper 3 expansion diagrams ===')
    what_evolves_p3_path = gen_what_evolves_p3()
    masbench_path = gen_masbench_axes()
    edge_comp_path = gen_edge_competence()

    print('\n=== Phase 4: Generating transition diagrams ===')
    trans_p1p2_path = gen_transition_p1_p2()
    trans_p2p3_path = gen_transition_p2_p3()

    print('\n=== Phase 5: Generating layout-fix diagrams ===')
    five_dec_path = gen_five_decisions_fixed()
    synth_path = gen_final_synthesis()

    print('\n=== Phase 5b: Cropping paper screenshots ===')
    crops = {
        'paper_eo_fig1.png': 0.22,
        'paper_eo_main_results.png': 0.05,
        'paper_eo_topology.png': 0.08,
        'paper_mz_ablation.png': 0.08,
        'paper_mz_main_results.png': 0.05,
        'paper_mo_paradigm.png': 0.12,
    }
    cropped_paths = {}
    for name, pct in crops.items():
        result = crop_paper_image(name, pct)
        if result:
            cropped_paths[name] = result

    # ── Now modify the PPTX ──
    print('\n=== Modifying PPTX ===')
    prs = Presentation(str(PPTX_PATH))
    slides = prs.slides

    # Phase 1: Replace challenge slide images
    print('Replacing challenge slides 6-8...')
    replace_slide_image(slides[5], c1_path)
    replace_slide_image(slides[6], c2_path)
    replace_slide_image(slides[7], c3_path)

    # Phase 2: Fix duplicates
    print('Fixing duplicate content...')
    replace_slide_image(slides[9], p1_motiv_path)  # Slide 10
    replace_slide_image(slides[22], meta_loop_path)  # Slide 23

    update_title_text(slides[11], 'What Evolves? (Puppeteer)')  # Slide 12
    update_title_text(slides[21], 'What Evolves? (MAS-ZERO)')  # Slide 22

    # Phase 5: Fix slide 3 text overflow
    print('Fixing slide 3 text overflow...')
    replace_slide_image(slides[2], five_dec_path)

    # Phase 5b: Replace cropped paper screenshots
    print('Replacing cropped paper screenshots...')
    crop_map = {
        14: 'paper_eo_fig1.png',       # Slide 15
        16: 'paper_eo_main_results.png',  # Slide 17
        17: 'paper_eo_topology.png',    # Slide 18
        27: 'paper_mz_ablation.png',    # Slide 28
        26: 'paper_mz_main_results.png',  # Slide 27
        30: 'paper_mo_paradigm.png',    # Slide 31
    }
    for idx, name in crop_map.items():
        if name in cropped_paths:
            replace_slide_image(slides[idx], cropped_paths[name])

    # Phase 6: Update final synthesis
    print('Updating final synthesis slide...')
    replace_slide_image(slides[34], synth_path)  # Slide 35

    # Phase 3 & 4: Insert new slides (reverse order to preserve indices)
    print('Inserting new slides...')

    # Slide after 34: Edge of Competence
    insert_slide_after(prs, 34, 'Key Insight: When MAS Helps Most',
                       edge_comp_path,
                       'MAS is most effective at the edge of sub-agent competence.')

    # Slide after 33: MASBench Five Axes
    insert_slide_after(prs, 33, 'MASBench: Five Axes of Complexity',
                       masbench_path,
                       'Controlled axes enable systematic MAS vs SAS comparison.',
                       'Source: MAS-Orchestra, arXiv:2601.14652')

    # Slide after 31: What Evolves MAS-Orchestra
    insert_slide_after(prs, 31, 'What Evolves? (MAS-Orchestra)',
                       what_evolves_p3_path,
                       'The full MAS structure is generated in one holistic decision.',
                       'Source: MAS-Orchestra, arXiv:2601.14652')

    # Slide after 29: Transition P2→P3
    insert_slide_after(prs, 29, 'From Self-Design to Understanding MAS',
                       trans_p2p3_path,
                       'Paper 2 removes supervision cost. Paper 3 asks: when does MAS actually help?')

    # Slide after 19: Transition P1→P2
    insert_slide_after(prs, 19, 'From Trained Control to Self-Design',
                       trans_p1p2_path,
                       'Paper 1 learns dynamic control via RL. Paper 2 removes the training requirement.')

    # Save
    out_path = PPTX_PATH
    prs.save(str(out_path))
    print(f'\nSaved to {out_path}')
    print(f'Total slides: {len(prs.slides)}')


if __name__ == '__main__':
    main()
